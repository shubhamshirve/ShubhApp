"""Operator router: profile, plans, subscribers, invoices, staff, reports, subscription, checkout, etc."""
from fastapi import APIRouter, HTTPException, Depends, Query, UploadFile, File, Request
from fastapi.responses import Response
from datetime import datetime, timezone, timedelta
from typing import List, Optional
import os
import csv
import io
import logging

from database import db
from models import (
    OperatorResponse, OperatorUpdate, InvoiceCustomization,
    AnnouncementCreate, OperatorPlanCreate, OperatorPlanResponse,
    SubscriberCreate, SubscriberResponse,
    InvoiceCreate, InvoiceResponse, PaymentLinkResponse,
    StaffCreate, StaffResponse, AuditLogResponse,
    PaymentGatewayConfig, SendNotificationRequest, BulkNotificationRequest,
    ReminderSettingsUpdate,
)
from utils import generate_id, hash_password, generate_invoice_number, generate_invoice_number_atomic
from dependencies import require_operator, require_operator_no_staff, check_operator_read_only
from audit import log_audit
from sanitization import sanitize_filename, sanitize_text

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/operator", tags=["Operator"])


# ── Addon helper ──────────────────────────────────────────────────────────────

async def _has_addon(operator_id: str, addon_code: str) -> bool:
    """Check if operator has an addon via their SaaS plan or purchased addons."""
    operator = await db.operators.find_one({"id": operator_id, "deleted_at": None}, {"_id": 0})
    if not operator:
        return False
    if addon_code in operator.get("active_addons", []):
        return True
    plan_id = operator.get("saas_plan_id")
    if plan_id:
        plan = await db.saas_plans.find_one({"id": plan_id, "deleted_at": None}, {"_id": 0})
        if plan and addon_code in plan.get("included_addons", []):
            return True
    return False


async def _get_platform_whatsapp_config():
    """Get the global platform WhatsApp config from admin settings."""
    config = await db.global_settings.find_one({"type": "platform_whatsapp"}, {"_id": 0})
    if not config or not config.get("access_token"):
        return None
    return config


async def _get_whatsapp_template_settings():
    """Get template assignment settings from admin config."""
    settings = await db.global_settings.find_one({"type": "whatsapp_template_settings"}, {"_id": 0})
    return settings or {}



# ─── Profile ─────────────────────────────────────────────────────────────────

def _parse_operator(o: dict) -> OperatorResponse:
    return OperatorResponse(**{
        **o,
        "created_at": datetime.fromisoformat(o["created_at"]),
        "trial_ends_at": datetime.fromisoformat(o["trial_ends_at"]) if o.get("trial_ends_at") else None,
        "subscription_ends_at": datetime.fromisoformat(o["subscription_ends_at"]) if o.get("subscription_ends_at") else None,
    })


# ─── Features endpoint ────────────────────────────────────────────────────────

@router.get("/features")
async def get_operator_features(current_user: dict = Depends(require_operator)):
    """Return which addon features are active for this operator."""
    if current_user["role"] == "admin":
        return {code: True for code in [
            "audit_log", "payment_gateway", "custom_payment_gateway",
            "announcement", "whatsapp_notifications",
            "staff_management"
        ]}
    operator_id = current_user["operator_id"]
    addon_codes = [
        "audit_log", "payment_gateway", "custom_payment_gateway",
        "announcement", "whatsapp_notifications",
        "staff_management"
    ]
    result = {code: await _has_addon(operator_id, code) for code in addon_codes}
    return result


@router.get("/profile", response_model=OperatorResponse)
async def get_operator_profile(current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin users don't have operator profile")
    operator = await db.operators.find_one({"id": current_user["operator_id"], "deleted_at": None}, {"_id": 0})
    if not operator:
        raise HTTPException(status_code=404, detail="Operator not found")
    return _parse_operator(operator)


@router.put("/profile", response_model=OperatorResponse)
async def update_operator_profile(data: OperatorUpdate, current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin users don't have operator profile")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode. Please renew subscription.")

    raw = data.model_dump()
    explicitly_set = data.model_fields_set  # fields actually sent in request body

    # Build update: only include fields that were explicitly sent
    update_data = {}
    for k, v in raw.items():
        if k in explicitly_set:
            update_data[k] = v  # could be None (for clearing), or a real value

    update_data.pop("status", None)

    # GST validation: cannot enable charge_gst without a valid GST number
    if update_data.get("charge_gst") is True:
        # Check if a valid gst_number will exist after this update
        new_gst = update_data.get("gst_number") if "gst_number" in update_data else None
        if new_gst:
            pass  # new valid GST provided in this update
        elif "gst_number" in update_data and not new_gst:
            # Explicitly clearing gst_number while enabling GST - not allowed
            raise HTTPException(
                status_code=400,
                detail="Cannot enable GST charging without a valid GSTIN. Please add your GST number first."
            )
        else:
            # gst_number not in update, check existing DB
            existing = await db.operators.find_one(
                {"id": current_user["operator_id"]}, {"_id": 0, "gst_number": 1}
            )
            if not (existing and existing.get("gst_number")):
                raise HTTPException(
                    status_code=400,
                    detail="Cannot enable GST charging without a valid GSTIN. Please add your GST number first."
                )

    update_data["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.operators.update_one({"id": current_user["operator_id"]}, {"$set": update_data})
    updated = await db.operators.find_one({"id": current_user["operator_id"]}, {"_id": 0})
    return _parse_operator(updated)


# ─── Invoice Settings ───────────────────────────────────────────────────────

@router.get("/invoice-settings")
async def get_invoice_settings(current_user: dict = Depends(require_operator)):
    settings = await db.invoice_settings.find_one({"operator_id": current_user["operator_id"]}, {"_id": 0})
    if not settings:
        operator = await db.operators.find_one({"id": current_user["operator_id"]}, {"_id": 0})
        return {
            "company_name": operator.get("company_name", ""),
            "company_address": "", "company_phone": operator.get("phone", ""),
            "company_email": operator.get("email", ""), "logo_url": None,
            "invoice_prefix": "INV", "invoice_footer": None, "show_gst": True, "terms_conditions": None,
            "invoice_template": "classic"  # Default value for new invoice template feature
        }
    
    # Ensure invoice_template field is present in existing records
    if "invoice_template" not in settings:
        settings["invoice_template"] = "classic"
    
    return settings


@router.put("/invoice-settings")
async def update_invoice_settings(data: InvoiceCustomization, current_user: dict = Depends(require_operator)):
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    now = datetime.now(timezone.utc)
    settings = {"operator_id": current_user["operator_id"], **data.model_dump(), "updated_at": now.isoformat()}
    await db.invoice_settings.update_one(
        {"operator_id": current_user["operator_id"]}, {"$set": settings}, upsert=True
    )
    return {"message": "Invoice settings updated"}


# ─── Theme Settings ─────────────────────────────────────────────────────────

@router.get("/theme-settings")
async def get_theme_settings(current_user: dict = Depends(require_operator)):
    """Get theme settings for the operator."""
    settings = await db.operator_theme.find_one(
        {"operator_id": current_user["operator_id"]}, {"_id": 0}
    )
    if not settings:
        # Return default based on plan
        operator = await db.operators.find_one({"id": current_user["operator_id"]}, {"_id": 0})
        is_pro = operator and operator.get("saas_plan_name", "").lower() in ["pro", "professional", "enterprise"]
        return {"theme": "classic" if is_pro else "modern", "operator_id": current_user["operator_id"]}
    return settings


@router.put("/theme-settings")
async def update_theme_settings(theme: str, current_user: dict = Depends(require_operator)):
    """Update theme settings for the operator."""
    theme = sanitize_text(theme)
    if theme not in ["modern", "classic"]:
        raise HTTPException(status_code=400, detail="Invalid theme. Choose 'modern' or 'classic'")
    now = datetime.now(timezone.utc)
    await db.operator_theme.update_one(
        {"operator_id": current_user["operator_id"]},
        {"$set": {"operator_id": current_user["operator_id"], "theme": theme, "updated_at": now.isoformat()}},
        upsert=True
    )
    return {"message": f"Theme updated to {theme}", "theme": theme}


# ─── Announcements ─────────────────────────────────────────────────────────

@router.post("/announcements")
async def create_announcement(data: AnnouncementCreate, current_user: dict = Depends(require_operator)):
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")

    # Check announcement addon
    if not await _has_addon(current_user["operator_id"], "announcement"):
        raise HTTPException(status_code=403, detail="Announcement add-on is not enabled for your plan.")

    # Enforce max 3 announcements per day
    today_start = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0).isoformat()
    today_count = await db.announcements.count_documents({
        "operator_id": current_user["operator_id"],
        "created_at": {"$gte": today_start}
    })
    if today_count >= 3:
        raise HTTPException(status_code=429, detail="Daily announcement limit reached (max 3 per day).")

    now = datetime.now(timezone.utc)
    if data.send_to_all:
        subscribers = await db.subscribers.find(
            {"operator_id": current_user["operator_id"], "status": "active", "deleted_at": None}, {"_id": 0}
        ).to_list(10000)
    else:
        subscribers = await db.subscribers.find(
            {"id": {"$in": data.subscriber_ids or []}, "operator_id": current_user["operator_id"], "deleted_at": None},
            {"_id": 0}
        ).to_list(10000)

    announcement = {
        "id": generate_id(), "operator_id": current_user["operator_id"],
        "title": data.title, "message": data.message,
        "recipient_count": len(subscribers), "sent_via_whatsapp": data.send_whatsapp,
        "created_by": current_user["id"], "created_at": now.isoformat()
    }
    await db.announcements.insert_one(announcement)
    announcement.pop("_id", None)

    sent_count = 0
    if data.send_whatsapp:
        for sub in subscribers:
            notification = {
                "id": generate_id(), "operator_id": current_user["operator_id"],
                "subscriber_id": sub["id"], "notification_type": "announcement",
                "whatsapp_number": sub["whatsapp_number"],
                "message": f"*{data.title}*\n\n{data.message}",
                "status": "pending", "created_at": now.isoformat()
            }
            await db.notification_queue.insert_one(notification)
            sent_count += 1

    return {"message": "Announcement created", "recipients": len(subscribers), "queued_notifications": sent_count}


@router.get("/announcements")
async def get_announcements(current_user: dict = Depends(require_operator)):
    announcements = await db.announcements.find(
        {"operator_id": current_user["operator_id"]}, {"_id": 0}
    ).sort("created_at", -1).to_list(100)
    return announcements


# ─── Addon Store & Subscription & Checkout ──────────────────────────────────────

@router.get("/addons/store")
async def get_addon_store(current_user: dict = Depends(require_operator)):
    # Block pure admin (not impersonating)
    if current_user["role"] == "admin" and not current_user.get("impersonated_by"):
        raise HTTPException(status_code=400, detail="Admin does not purchase addons")
    operator = await db.operators.find_one({"id": current_user["operator_id"], "deleted_at": None}, {"_id": 0})
    if not operator:
        raise HTTPException(status_code=404, detail="Operator not found")
    all_addons = await db.addons.find({"deleted_at": None}, {"_id": 0}).to_list(100)
    active = operator.get("active_addons", [])
    addon_expiry = operator.get("addon_expiry", {})
    plan_addons = []
    if operator.get("saas_plan_id"):
        plan = await db.saas_plans.find_one({"id": operator["saas_plan_id"], "deleted_at": None}, {"_id": 0})
        if plan:
            plan_addons = plan.get("included_addons", [])
    result = []
    for addon in all_addons:
        status = "available"
        expires_at = None
        if addon["code"] in active:
            status = "purchased"
            expires_at = addon_expiry.get(addon["code"])
        elif addon["code"] in plan_addons:
            status = "included_in_plan"
            expires_at = operator.get("subscription_ends_at")
        result.append({
            "id": addon["id"], "name": addon["name"], "code": addon["code"],
            "price": addon["price"], "description": addon.get("description", ""),
            "status": status, "expires_at": expires_at
        })
    return result


@router.post("/checkout/validate-coupon")
async def validate_coupon(
    code: str, amount: float,
    current_user: dict = Depends(require_operator)
):
    """Validate a discount code and return the discount amount."""
    if current_user["role"] == "admin" and not current_user.get("impersonated_by"):
        raise HTTPException(status_code=400, detail="Admin cannot use coupons")
    now = datetime.now(timezone.utc)
    doc = await db.discount_codes.find_one({"code": code.upper(), "deleted_at": None}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Invalid discount code")
    if not doc.get("is_active", True):
        raise HTTPException(status_code=400, detail="Discount code is inactive")
    if doc.get("expiry_date"):
        expiry = datetime.fromisoformat(doc["expiry_date"])
        if expiry.tzinfo is None:
            expiry = expiry.replace(tzinfo=timezone.utc)
        if now > expiry:
            raise HTTPException(status_code=400, detail="Discount code has expired")
    max_r = doc.get("max_redemptions", 0)
    if max_r > 0 and doc.get("used_count", 0) >= max_r:
        raise HTTPException(status_code=400, detail="Discount code redemption limit reached")
    # Per-operator: each operator can only use a coupon once in their lifetime
    operator_id = current_user.get("operator_id", "")
    if operator_id and operator_id in doc.get("redeemed_by", []):
        raise HTTPException(status_code=400, detail="You have already used this discount code")
    # Calculate discount
    if doc["discount_type"] == "percentage":
        discount_amount = round(amount * doc["discount_value"] / 100, 2)
    else:
        discount_amount = min(doc["discount_value"], amount)
    final_amount = max(0, amount - discount_amount)
    return {
        "valid": True,
        "code": doc["code"],
        "discount_type": doc["discount_type"],
        "discount_value": doc["discount_value"],
        "discount_amount": discount_amount,
        "final_amount": final_amount,
        "message": f"{'{}%'.format(int(doc['discount_value'])) if doc['discount_type'] == 'percentage' else '₹{}'.format(int(doc['discount_value']))} discount applied!"
    }


@router.post("/checkout/create-order")
async def create_checkout_order(
    item_type: str, item_code: str = "", months: int = 1, plan_id: str = "",
    addon_codes: str = "", coupon_code: str = "",
    current_user: dict = Depends(require_operator)
):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot checkout")
    operator = await db.operators.find_one({"id": current_user["operator_id"], "deleted_at": None}, {"_id": 0})
    if not operator:
        raise HTTPException(status_code=404, detail="Operator not found")
    settings = await db.global_settings.find_one({"type": "platform"}, {"_id": 0})
    gst_rate = settings.get("gst_rate", 18) if settings else 18
    description = base_amount = receipt_prefix = ""
    base_amount = 0
    selected_addon_codes = [c.strip() for c in addon_codes.split(",") if c.strip()] if addon_codes else []

    if item_type == "addon":
        # Block addon purchase on trial plan
        if operator.get("status") == "trial":
            raise HTTPException(status_code=403, detail="Please subscribe to a paid plan to purchase add-ons.")
        addon = await db.addons.find_one({"code": item_code, "deleted_at": None}, {"_id": 0})
        if not addon:
            raise HTTPException(status_code=404, detail="Add-on not found")
        if item_code in operator.get("active_addons", []):
            raise HTTPException(status_code=400, detail="Add-on already active")
        if operator.get("saas_plan_id"):
            sp = await db.saas_plans.find_one({"id": operator["saas_plan_id"], "deleted_at": None}, {"_id": 0})
            if sp and item_code in sp.get("included_addons", []):
                active = operator.get("active_addons", [])
                active.append(item_code)
                await db.operators.update_one({"id": operator["id"]}, {"$set": {"active_addons": active}})
                return {"status": "activated_free", "message": f"{addon['name']} is included in your plan and activated!"}
        base_amount = addon["price"]
        description = f"Add-on: {addon['name']} (Monthly)"
        receipt_prefix = "ADDON"

    elif item_type == "subscription":
        target_plan_id = plan_id or operator.get("saas_plan_id")
        if not target_plan_id:
            raise HTTPException(status_code=400, detail="No plan selected")
        saas_plan = await db.saas_plans.find_one({"id": target_plan_id, "deleted_at": None}, {"_id": 0})
        if not saas_plan:
            raise HTTPException(status_code=404, detail="Plan not found")
        base_amount = saas_plan["monthly_price"] * months
        description = f"{saas_plan['name']} x {months} month(s)"
        receipt_prefix = "SUB"
        included_in_plan = saas_plan.get("included_addons", [])
        # Auto-include prices of operator's standalone purchased addons (not in plan)
        auto_addon_codes = []
        for code in operator.get("active_addons", []):
            if code not in included_in_plan:
                addon_doc = await db.addons.find_one({"code": code, "deleted_at": None}, {"_id": 0})
                if addon_doc:
                    base_amount += addon_doc["price"] * months
                    auto_addon_codes.append(code)
        # Add newly selected addon prices (not already owned)
        valid_addon_codes = []
        for code in selected_addon_codes:
            if code in auto_addon_codes:
                continue  # Already counted
            addon_doc2 = await db.addons.find_one({"code": code, "deleted_at": None}, {"_id": 0})
            if addon_doc2 and code not in operator.get("active_addons", []) and code not in included_in_plan:
                base_amount += addon_doc2["price"] * months
                valid_addon_codes.append(code)
        all_addon_codes = auto_addon_codes + valid_addon_codes
        if all_addon_codes:
            description += f" + {len(all_addon_codes)} add-on(s)"
        selected_addon_codes = all_addon_codes
    else:
        raise HTTPException(status_code=400, detail="Invalid item_type. Use 'addon' or 'subscription'")

    if base_amount <= 0:
        raise HTTPException(status_code=400, detail="Amount must be greater than zero")

    # Apply coupon discount (before GST)
    discount_amount = 0.0
    applied_coupon = None
    if coupon_code:
        coupon_doc = await db.discount_codes.find_one({"code": coupon_code.upper(), "deleted_at": None}, {"_id": 0})
        if coupon_doc and coupon_doc.get("is_active"):
            now_check = datetime.now(timezone.utc)
            valid = True
            if coupon_doc.get("expiry_date"):
                exp = datetime.fromisoformat(coupon_doc["expiry_date"])
                if exp.tzinfo is None:
                    exp = exp.replace(tzinfo=timezone.utc)
                if now_check > exp:
                    valid = False
            max_r = coupon_doc.get("max_redemptions", 0)
            if max_r > 0 and coupon_doc.get("used_count", 0) >= max_r:
                valid = False
            if valid:
                # Per-operator: each operator can only use a coupon once in their lifetime
                if operator["id"] in coupon_doc.get("redeemed_by", []):
                    valid = False
            if valid:
                if coupon_doc["discount_type"] == "percentage":
                    discount_amount = round(base_amount * coupon_doc["discount_value"] / 100, 2)
                else:
                    discount_amount = min(coupon_doc["discount_value"], base_amount)
                applied_coupon = coupon_code.upper()
    discounted_base = round(base_amount - discount_amount, 2)

    gst_amount = round(discounted_base * gst_rate / 100, 2)
    exact_total = round(discounted_base + gst_amount, 2)
    import math
    rounded_total = math.floor(exact_total + 0.5)          # standard half-up rounding → int
    rounding_diff = round(rounded_total - exact_total, 2)  # +ve = rounded up, -ve = rounded down

    # Use platform payment gateway from DB, fallback to env vars
    platform_gw = await db.payment_gateways.find_one({"is_platform_gateway": True, "is_active": True}, {"_id": 0})
    if platform_gw:
        razorpay_key = platform_gw["api_key"]
        razorpay_secret = platform_gw["api_secret"]
    else:
        razorpay_key = os.environ.get("RAZORPAY_KEY_ID")
        razorpay_secret = os.environ.get("RAZORPAY_KEY_SECRET")
    if not razorpay_key or not razorpay_secret:
        raise HTTPException(status_code=500, detail="Platform payment gateway not configured. Please contact admin.")

    from services.razorpay_service import RazorpayService
    rz = RazorpayService(razorpay_key, razorpay_secret)
    order_id = generate_id()
    order = rz.create_order(
        amount=rounded_total, receipt=f"{receipt_prefix}-{order_id[:8]}",
        notes={"type": item_type, "item_code": item_code, "plan_id": plan_id,
               "months": str(months), "operator_id": operator["id"], "internal_order_id": order_id}
    )

    now = datetime.now(timezone.utc)
    await db.checkout_orders.insert_one({
        "id": order_id, "razorpay_order_id": order["id"],
        "operator_id": operator["id"], "item_type": item_type,
        "item_code": item_code, "plan_id": plan_id, "months": months,
        "addon_codes": selected_addon_codes,
        "base_amount": base_amount, "discount_amount": discount_amount,
        "discounted_base": discounted_base,
        "gst_amount": gst_amount,
        "exact_total": exact_total, "rounding_diff": rounding_diff,
        "total_amount": rounded_total,
        "coupon_code": applied_coupon,
        "description": description, "status": "created",
        "created_at": now.isoformat(), "deleted_at": None
    })

    platform_name = settings.get("platform_name", "E-Bill") if settings else "E-Bill"
    return {
        "razorpay_order_id": order["id"], "razorpay_key": razorpay_key,
        "amount": rounded_total * 100, "currency": "INR",
        "name": platform_name, "description": description,
        "base_amount": base_amount, "discount_amount": discount_amount,
        "discounted_base": discounted_base,
        "gst_amount": gst_amount,
        "exact_total": exact_total, "rounding_diff": rounding_diff,
        "total_amount": rounded_total,
        "coupon_code": applied_coupon,
        "prefill": {"name": operator.get("owner_name", ""), "email": operator.get("email", ""), "contact": operator.get("phone", "")}
    }


@router.post("/checkout/verify")
async def verify_checkout_payment(
    razorpay_order_id: str, razorpay_payment_id: str, razorpay_signature: str,
    current_user: dict = Depends(require_operator)
):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot verify checkout")
    platform_gw = await db.payment_gateways.find_one({"is_platform_gateway": True, "is_active": True}, {"_id": 0})
    if platform_gw:
        razorpay_key = platform_gw["api_key"]
        razorpay_secret = platform_gw["api_secret"]
    else:
        razorpay_key = os.environ.get("RAZORPAY_KEY_ID")
        razorpay_secret = os.environ.get("RAZORPAY_KEY_SECRET")
    if not razorpay_key or not razorpay_secret:
        raise HTTPException(status_code=500, detail="Platform payment gateway not configured. Please contact admin.")

    from services.razorpay_service import RazorpayService
    rz = RazorpayService(razorpay_key, razorpay_secret)
    if not rz.verify_payment_signature(razorpay_order_id, razorpay_payment_id, razorpay_signature):
        raise HTTPException(status_code=400, detail="Payment verification failed")

    order = await db.checkout_orders.find_one(
        {"razorpay_order_id": razorpay_order_id, "operator_id": current_user["operator_id"]}, {"_id": 0}
    )
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    now = datetime.now(timezone.utc)
    operator = await db.operators.find_one({"id": current_user["operator_id"]}, {"_id": 0})

    if order["item_type"] == "addon":
        active = operator.get("active_addons", [])
        addon_expiry = operator.get("addon_expiry", {})
        new_code = order["item_code"]
        # Mutual exclusion: payment_gateway and custom_payment_gateway cannot coexist
        if new_code == "payment_gateway" and "custom_payment_gateway" in active:
            active.remove("custom_payment_gateway")
            addon_expiry.pop("custom_payment_gateway", None)
        elif new_code == "custom_payment_gateway" and "payment_gateway" in active:
            active.remove("payment_gateway")
            addon_expiry.pop("payment_gateway", None)
        if new_code not in active:
            active.append(new_code)
        # Expiry = current subscription end date
        expiry_date = operator.get("subscription_ends_at") or now.isoformat()
        addon_expiry[new_code] = expiry_date
        await db.operators.update_one(
            {"id": operator["id"]},
            {"$set": {"active_addons": active, "addon_expiry": addon_expiry, "updated_at": now.isoformat()}}
        )
        result_msg = f"Add-on '{new_code}' activated"

    elif order["item_type"] == "subscription":
        from dateutil.relativedelta import relativedelta
        current_end = operator.get("subscription_ends_at")
        start = now
        if current_end:
            start = datetime.fromisoformat(current_end)
            if start < now:
                start = now
        new_end = start + relativedelta(months=order["months"])
        new_end_iso = new_end.isoformat()
        update_fields = {
            "subscription_ends_at": new_end_iso, "status": "active",
            "is_read_only": False, "updated_at": now.isoformat()
        }
        if order.get("plan_id"):
            update_fields["saas_plan_id"] = order["plan_id"]
        # Extend all existing active addons expiry to new subscription end date
        addon_expiry = operator.get("addon_expiry", {})
        active = operator.get("active_addons", [])
        for code in active:
            addon_expiry[code] = new_end_iso
        # Activate any addons bundled with this subscription order
        if order.get("addon_codes"):
            for code in order["addon_codes"]:
                # Mutual exclusion check for payment gateway addons
                if code == "payment_gateway" and "custom_payment_gateway" in active:
                    active.remove("custom_payment_gateway")
                    addon_expiry.pop("custom_payment_gateway", None)
                elif code == "custom_payment_gateway" and "payment_gateway" in active:
                    active.remove("payment_gateway")
                    addon_expiry.pop("payment_gateway", None)
                if code not in active:
                    active.append(code)
                    # If staff_management addon, set max_staff = 5
                    if code == "staff_management":
                        update_fields["max_staff"] = 5
                addon_expiry[code] = new_end_iso
            update_fields["active_addons"] = active
        update_fields["addon_expiry"] = addon_expiry
        await db.operators.update_one({"id": operator["id"]}, {"$set": update_fields})
        result_msg = f"Subscription extended by {order['months']} month(s)"
    else:
        result_msg = "Payment verified"

    await db.checkout_orders.update_one(
        {"razorpay_order_id": razorpay_order_id},
        {"$set": {"status": "paid", "razorpay_payment_id": razorpay_payment_id, "paid_at": now.isoformat()}}
    )

    # Increment coupon used_count and record operator if a coupon was applied
    if order.get("coupon_code"):
        await db.discount_codes.update_one(
            {"code": order["coupon_code"]},
            {
                "$inc": {"used_count": 1},
                "$addToSet": {"redeemed_by": operator["id"]}
            }
        )

    payment_record = {
        "id": generate_id(), "operator_id": operator["id"],
        "order_id": order["id"], "razorpay_order_id": razorpay_order_id,
        "razorpay_payment_id": razorpay_payment_id,
        "item_type": order["item_type"], "item_code": order.get("item_code", ""),
        "base_amount": order["base_amount"],
        "discount_amount": order.get("discount_amount", 0),
        "coupon_code": order.get("coupon_code"),
        "gst_amount": order["gst_amount"],
        "total_amount": order["total_amount"], "status": "completed",
        "created_at": now.isoformat(), "deleted_at": None
    }
    await db.saas_payments.insert_one(payment_record)
    await log_audit(
        current_user["id"], current_user["name"], current_user["role"],
        "payment", "checkout", None, {"type": order["item_type"], "amount": order["total_amount"]},
        ip_address=current_user.get("_ip_address"),
        operator_id=operator["id"]
    )
    return {"status": "success", "message": result_msg}


@router.get("/subscription")
async def get_operator_subscription(current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin does not have a subscription")
    operator = await db.operators.find_one({"id": current_user["operator_id"], "deleted_at": None}, {"_id": 0})
    if not operator:
        raise HTTPException(status_code=404, detail="Operator not found")
    saas_plan = None
    if operator.get("saas_plan_id"):
        saas_plan = await db.saas_plans.find_one({"id": operator["saas_plan_id"], "deleted_at": None}, {"_id": 0})
    available_plans = await db.saas_plans.find({"deleted_at": None, "trial_enabled": False}, {"_id": 0}).to_list(50)
    return {
        "operator_id": operator["id"], "company_name": operator.get("company_name", ""),
        "status": operator.get("status", "unknown"),
        "saas_plan_id": operator.get("saas_plan_id"),
        "saas_plan_name": operator.get("saas_plan_name") or (saas_plan["name"] if saas_plan else None),
        "saas_plan_price": saas_plan.get("monthly_price") if saas_plan else None,
        "subscription_ends_at": operator.get("subscription_ends_at"),
        "trial_ends_at": operator.get("trial_ends_at"),
        "is_read_only": operator.get("is_read_only", False),
        "available_plans": [{"id": p["id"], "name": p["name"], "monthly_price": p["monthly_price"]} for p in available_plans]
    }


@router.get("/payment-history")
async def get_operator_payment_history(current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin does not have payment history")
    payments = await db.saas_payments.find(
        {"operator_id": current_user["operator_id"], "status": "completed", "deleted_at": None}, {"_id": 0}
    ).sort("created_at", -1).to_list(200)
    return payments


@router.post("/renew-subscription")
async def renew_operator_subscription(
    plan_id: Optional[str] = None, months: int = 1,
    current_user: dict = Depends(require_operator)
):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot renew")
    operator = await db.operators.find_one({"id": current_user["operator_id"], "deleted_at": None}, {"_id": 0})
    if not operator:
        raise HTTPException(status_code=404, detail="Operator not found")
    target_plan_id = plan_id or operator.get("saas_plan_id")
    if not target_plan_id:
        raise HTTPException(status_code=400, detail="No plan selected.")
    saas_plan = await db.saas_plans.find_one({"id": target_plan_id, "deleted_at": None}, {"_id": 0})
    if not saas_plan:
        raise HTTPException(status_code=404, detail="Plan not found")
    amount = saas_plan["monthly_price"] * months
    gst_amount = 0
    if saas_plan.get("gst_applicable"):
        settings = await db.global_settings.find_one({"type": "platform"}, {"_id": 0})
        gst_rate = settings.get("gst_rate", 18) if settings else 18
        gst_amount = round(amount * gst_rate / 100, 2)
    import math
    exact_total = round(amount + gst_amount, 2)
    total_amount = math.floor(exact_total + 0.5)           # standard half-up round to int
    rounding_diff = round(total_amount - exact_total, 2)
    renewal_id = generate_id()
    payment_link = None
    razorpay_key = os.environ.get("RAZORPAY_KEY_ID")
    razorpay_secret = os.environ.get("RAZORPAY_KEY_SECRET")
    if razorpay_key and razorpay_secret and total_amount > 0:
        try:
            from services.razorpay_service import RazorpayService
            rz = RazorpayService(razorpay_key, razorpay_secret)
            result = rz.create_payment_link(
                amount=total_amount,
                description=f"SaaS Subscription: {saas_plan['name']} x {months} month(s)",
                customer_name=operator.get("owner_name", "Operator"),
                customer_email=operator.get("email", ""),
                customer_phone=operator.get("phone", ""),
                invoice_number=f"SAAS-{renewal_id[:8]}"
            )
            if result:
                payment_link = result.get("short_url")
        except Exception as e:
            logger.warning(f"Razorpay link creation failed: {e}")
    now = datetime.now(timezone.utc)
    renewal = {
        "id": renewal_id, "operator_id": operator["id"],
        "plan_id": target_plan_id, "plan_name": saas_plan["name"],
        "months": months, "base_amount": amount, "gst_amount": gst_amount,
        "exact_total": exact_total, "rounding_diff": rounding_diff,
        "total_amount": total_amount, "payment_link": payment_link,
        "status": "pending", "created_at": now.isoformat(), "deleted_at": None
    }
    await db.saas_payments.insert_one(renewal)
    renewal.pop("_id", None)
    return renewal


# ─── Dashboard ──────────────────────────────────────────────────────────────

@router.get("/dashboard")
async def get_operator_dashboard(current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Use admin dashboard")
    operator_id = current_user["operator_id"]
    total_subscribers = await db.subscribers.count_documents({"operator_id": operator_id, "deleted_at": None})
    active_subscribers = await db.subscribers.count_documents({"operator_id": operator_id, "status": "active", "deleted_at": None})
    total_invoices = await db.invoices.count_documents({"operator_id": operator_id, "deleted_at": None})
    pending_invoices = await db.invoices.count_documents({"operator_id": operator_id, "status": "pending", "deleted_at": None})
    overdue_invoices = await db.invoices.count_documents({"operator_id": operator_id, "status": "overdue", "deleted_at": None})
    paid_invoices = await db.invoices.count_documents({"operator_id": operator_id, "status": "paid", "deleted_at": None})
    paid_invoice_list = await db.invoices.find(
        {"operator_id": operator_id, "status": "paid", "deleted_at": None}, {"_id": 0, "final_amount": 1}
    ).to_list(1000)
    total_revenue = sum(inv.get("final_amount", 0) for inv in paid_invoice_list)
    operator = await db.operators.find_one({"id": operator_id, "deleted_at": None}, {"_id": 0})
    # Fetch plan limits
    max_subscribers = None
    max_staff = None
    if operator and operator.get("saas_plan_id"):
        sp = await db.saas_plans.find_one({"id": operator["saas_plan_id"], "deleted_at": None}, {"_id": 0})
        if sp:
            max_subscribers = sp.get("max_subscribers")
            # Use operator-level max_staff override if set (staff_management addon)
            max_staff = operator.get("max_staff") or sp.get("max_staff")
    current_staff = await db.users.count_documents({"operator_id": operator_id, "role": "staff", "deleted_at": None})
    slots_remaining = max(0, max_subscribers - total_subscribers) if max_subscribers is not None else None
    return {
        "total_subscribers": total_subscribers, "active_subscribers": active_subscribers,
        "total_invoices": total_invoices, "pending_invoices": pending_invoices,
        "overdue_invoices": overdue_invoices, "paid_invoices": paid_invoices,
        "total_revenue": total_revenue,
        "is_read_only": operator.get("is_read_only", False) if operator else False,
        "subscription_ends_at": operator.get("subscription_ends_at") if operator else None,
        "trial_ends_at": operator.get("trial_ends_at") if operator else None,
        "status": operator.get("status") if operator else "unknown",
        "max_subscribers": max_subscribers,
        "subscriber_slots_remaining": slots_remaining,
        "max_staff": max_staff,
        "current_staff": current_staff,
    }


# ─── Service Plans ──────────────────────────────────────────────────────────

@router.post("/plans", response_model=OperatorPlanResponse)
async def create_operator_plan(data: OperatorPlanCreate, current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot create operator plans")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    now = datetime.now(timezone.utc)
    plan = {
        "id": generate_id(), "name": data.name, "price": data.price, "validity": data.validity,
        "tax_percentage": data.tax_percentage, "tax_type": data.tax_type,
        "description": data.description, "status": "active",
        "operator_id": current_user["operator_id"],
        "created_at": now.isoformat(), "updated_at": now.isoformat(), "deleted_at": None
    }
    await db.operator_plans.insert_one(plan)
    return OperatorPlanResponse(**{**plan, "created_at": now})


@router.get("/plans", response_model=List[OperatorPlanResponse])
async def get_operator_plans(current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot access operator plans")
    try:
        plans = await db.operator_plans.find(
            {"operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
        ).to_list(100)
        parsed_plans = []
        for p in plans:
            try:
                created_at = p["created_at"]
                # Handle both string and datetime formats
                if isinstance(created_at, str):
                    created_at = datetime.fromisoformat(created_at)
                elif not isinstance(created_at, datetime):
                    created_at = datetime.now(timezone.utc)
                parsed_plans.append(OperatorPlanResponse(**{**p, "created_at": created_at}))
            except Exception as e:
                logger.error(f"Error parsing plan {p.get('id', 'unknown')}: {e}")
                # Skip this plan if parsing fails
                continue
        return parsed_plans
    except Exception as e:
        logger.error(f"Error in get_operator_plans: {e}")
        raise HTTPException(status_code=500, detail=f"Error retrieving plans: {str(e)}")


@router.put("/plans/{plan_id}", response_model=OperatorPlanResponse)
async def update_operator_plan(plan_id: str, data: OperatorPlanCreate, current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot update operator plans")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    existing = await db.operator_plans.find_one(
        {"id": plan_id, "operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
    )
    if not existing:
        raise HTTPException(status_code=404, detail="Plan not found")
    update_data = data.model_dump()
    update_data["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.operator_plans.update_one({"id": plan_id}, {"$set": update_data})
    updated = await db.operator_plans.find_one({"id": plan_id}, {"_id": 0})
    return OperatorPlanResponse(**{**updated, "created_at": datetime.fromisoformat(updated["created_at"])})


@router.delete("/plans/{plan_id}")
async def delete_operator_plan(plan_id: str, current_user: dict = Depends(require_operator_no_staff)):
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    result = await db.operator_plans.update_one(
        {"id": plan_id, "operator_id": current_user["operator_id"], "deleted_at": None},
        {"$set": {"deleted_at": datetime.now(timezone.utc).isoformat()}}
    )
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Plan not found")
    return {"message": "Plan deleted"}


# ─── Bulk Upload: Plans ────────────────────────────────────────────────────────

@router.get("/plans/sample-csv")
async def get_plans_sample_csv(current_user: dict = Depends(require_operator)):
    """Download a sample CSV template for bulk plan upload."""
    rows = [
        ["name", "price", "validity", "tax_percentage", "tax_type", "description"],
        ["Monthly Basic",      "500",  "monthly",     "18", "exclusive", "Basic monthly broadband plan"],
        ["Quarterly Standard", "1400", "quarterly",   "18", "exclusive", "Standard quarterly plan"],
        ["Half Yearly Gold",   "2700", "half_yearly",  "0", "none",      "Half yearly plan with no tax"],
        ["Annual Premium",     "5000", "yearly",      "18", "inclusive", "Annual premium plan GST inclusive"],
    ]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerows(rows)
    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=plans_sample.csv"}
    )


@router.post("/plans/bulk-upload")
async def bulk_upload_plans(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_operator)
):
    """Bulk upload service plans from a CSV or XLSX file."""
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot upload operator plans")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")

    content = await file.read()
    filename = sanitize_filename(file.filename or "", default="plans").lower()
    rows = []
    VALID_VALIDITY = ["monthly", "quarterly", "half_yearly", "yearly"]
    VALID_TAX_TYPE = ["inclusive", "exclusive", "none"]

    try:
        if filename.endswith(".xlsx") or filename.endswith(".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb.active
            headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append({headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
        else:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            for row in reader:
                rows.append({k.strip().lower(): v.strip() for k, v in row.items()})
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse file: {e}")

    now = datetime.now(timezone.utc)
    created, skipped, errors = [], [], []

    for idx, row in enumerate(rows, start=2):
        name = row.get("name", "").strip()
        if not name:
            errors.append({"row": idx, "reason": "name is required"})
            continue

        try:
            price = float(row.get("price", 0) or 0)
        except ValueError:
            errors.append({"row": idx, "name": name, "reason": "Invalid price"})
            continue

        validity = row.get("validity", "monthly").strip().lower()
        if validity not in VALID_VALIDITY:
            errors.append({"row": idx, "name": name, "reason": f"Invalid validity '{validity}'. Use: {VALID_VALIDITY}"})
            continue

        tax_type = row.get("tax_type", "none").strip().lower()
        if tax_type not in VALID_TAX_TYPE:
            tax_type = "none"

        try:
            tax_percentage = float(row.get("tax_percentage", 0) or 0)
        except ValueError:
            tax_percentage = 0

        # Check for duplicate name
        dup = await db.operator_plans.find_one(
            {"name": name, "operator_id": current_user["operator_id"], "deleted_at": None}
        )
        if dup:
            skipped.append({"row": idx, "name": name, "reason": "Plan name already exists"})
            continue

        plan = {
            "id": generate_id(), "name": name, "price": price, "validity": validity,
            "tax_percentage": tax_percentage, "tax_type": tax_type,
            "description": row.get("description", "") or None,
            "status": "active", "operator_id": current_user["operator_id"],
            "created_at": now.isoformat(), "updated_at": now.isoformat(), "deleted_at": None
        }
        await db.operator_plans.insert_one(plan)
        created.append(name)

    return {
        "message": f"Bulk upload complete: {len(created)} created, {len(skipped)} skipped, {len(errors)} errors",
        "created": len(created), "skipped": len(skipped), "errors": errors[:20]
    }


# ─── Subscribers ────────────────────────────────────────────────────────────

@router.post("/subscribers", response_model=SubscriberResponse)
async def create_subscriber(data: SubscriberCreate, current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot create subscribers")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    operator = await db.operators.find_one({"id": current_user["operator_id"], "deleted_at": None}, {"_id": 0})
    if operator:
        plan = await db.saas_plans.find_one({"id": operator.get("saas_plan_id"), "deleted_at": None}, {"_id": 0})
        if plan:
            current_count = await db.subscribers.count_documents({"operator_id": current_user["operator_id"], "deleted_at": None})
            if current_count >= plan["max_subscribers"]:
                raise HTTPException(
                    status_code=403,
                    detail=f"Subscriber limit reached ({current_count}/{plan['max_subscribers']}). Please upgrade your plan to add more subscribers."
                )
    op_plan = await db.operator_plans.find_one(
        {"id": data.plan_id, "operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
    )
    if not op_plan:
        raise HTTPException(status_code=404, detail="Plan not found")
    now = datetime.now(timezone.utc)
    subscriber = {
        "id": generate_id(), "name": data.name, "whatsapp_number": data.whatsapp_number,
        "email": data.email, "address": data.address, "plan_id": data.plan_id,
        "plan_name": op_plan["name"], "billing_date": data.billing_date, "discount": data.discount,
        "status": "active", "operator_id": current_user["operator_id"],
        "created_at": now.isoformat(), "updated_at": now.isoformat(), "deleted_at": None
    }
    await db.subscribers.insert_one(subscriber)
    return SubscriberResponse(**{**subscriber, "created_at": now})


@router.get("/subscribers", response_model=List[SubscriberResponse])
async def get_subscribers(
    status: Optional[str] = None, plan_id: Optional[str] = None,
    current_user: dict = Depends(require_operator)
):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot access operator subscribers")
    query = {"operator_id": current_user["operator_id"], "deleted_at": None}
    if status:
        query["status"] = status
    if plan_id:
        query["plan_id"] = plan_id
    subscribers = await db.subscribers.find(query, {"_id": 0}).to_list(1000)
    return [SubscriberResponse(**{**s, "created_at": datetime.fromisoformat(s["created_at"])}) for s in subscribers]


@router.get("/subscribers/{subscriber_id}", response_model=SubscriberResponse)
async def get_subscriber(subscriber_id: str, current_user: dict = Depends(require_operator)):
    subscriber = await db.subscribers.find_one(
        {"id": subscriber_id, "operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
    )
    if not subscriber:
        raise HTTPException(status_code=404, detail="Subscriber not found")
    return SubscriberResponse(**{**subscriber, "created_at": datetime.fromisoformat(subscriber["created_at"])})


@router.put("/subscribers/{subscriber_id}", response_model=SubscriberResponse)
async def update_subscriber(subscriber_id: str, data: SubscriberCreate, current_user: dict = Depends(require_operator)):
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    existing = await db.subscribers.find_one(
        {"id": subscriber_id, "operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
    )
    if not existing:
        raise HTTPException(status_code=404, detail="Subscriber not found")
    op_plan = await db.operator_plans.find_one({"id": data.plan_id, "deleted_at": None}, {"_id": 0})
    update_data = data.model_dump()
    update_data["plan_name"] = op_plan["name"] if op_plan else None
    update_data["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.subscribers.update_one({"id": subscriber_id}, {"$set": update_data})
    updated = await db.subscribers.find_one({"id": subscriber_id}, {"_id": 0})
    return SubscriberResponse(**{**updated, "created_at": datetime.fromisoformat(updated["created_at"])})


@router.delete("/subscribers/{subscriber_id}")
async def delete_subscriber(subscriber_id: str, current_user: dict = Depends(require_operator_no_staff)):
    # Only admin impersonating as operator can delete subscribers
    if not current_user.get("impersonated_by"):
        raise HTTPException(status_code=403, detail="Only admin can delete subscribers. Use suspend instead.")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    
    # Check for linked invoices that are not cancelled
    linked_invoices = await db.invoices.find(
        {"subscriber_id": subscriber_id, "operator_id": current_user["operator_id"], "deleted_at": None, "status": {"$ne": "cancelled"}}
    ).to_list(None)
    
    if linked_invoices:
        invoice_numbers = ", ".join([inv.get("invoice_number", inv.get("id", "")) for inv in linked_invoices[:5]])
        raise HTTPException(
            status_code=400, 
            detail=f"Cannot delete subscriber with active invoices. Please cancel these invoices first: {invoice_numbers}"
        )
    
    result = await db.subscribers.update_one(
        {"id": subscriber_id, "operator_id": current_user["operator_id"], "deleted_at": None},
        {"$set": {"deleted_at": datetime.now(timezone.utc).isoformat()}}
    )
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Subscriber not found")
    return {"message": "Subscriber deleted"}


@router.post("/subscribers/{subscriber_id}/suspend")
async def suspend_subscriber(subscriber_id: str, current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot suspend subscribers directly")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    result = await db.subscribers.update_one(
        {"id": subscriber_id, "operator_id": current_user["operator_id"], "deleted_at": None},
        {"$set": {"status": "inactive", "updated_at": datetime.now(timezone.utc).isoformat()}}
    )
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Subscriber not found")
    return {"message": "Subscriber suspended"}


@router.post("/subscribers/{subscriber_id}/activate")
async def activate_subscriber(subscriber_id: str, current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot activate subscribers directly")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    result = await db.subscribers.update_one(
        {"id": subscriber_id, "operator_id": current_user["operator_id"], "deleted_at": None},
        {"$set": {"status": "active", "updated_at": datetime.now(timezone.utc).isoformat()}}
    )
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Subscriber not found")
    return {"message": "Subscriber activated"}


# ─── Bulk Upload: Subscribers ─────────────────────────────────────────────────

@router.get("/subscribers/sample-csv")
async def get_subscribers_sample_csv(current_user: dict = Depends(require_operator)):
    """Download a sample CSV template for bulk subscriber upload."""
    rows = [
        ["name", "whatsapp_number", "email", "address", "plan_name", "billing_date", "discount"],
        ["Rajesh Kumar",   "9876543210", "rajesh@example.com",   "123 MG Road, Mumbai",    "Monthly Basic", "1",  "0"],
        ["Priya Sharma",   "9123456789", "priya@example.com",    "456 Anna Salai, Chennai", "Monthly Basic", "5",  "0"],
        ["Amit Patel",     "9988776655", "amit@example.com",     "789 FC Road, Pune",       "Monthly Basic", "10", "50"],
        ["Sunita Verma",   "9871234567", "",                     "321 Brigade Rd, Bangalore","Monthly Basic", "15", "0"],
    ]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerows(rows)
    csv_content = output.getvalue()
    return Response(
        content=csv_content,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=subscribers_sample.csv"}
    )


@router.post("/subscribers/bulk-upload")
async def bulk_upload_subscribers(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_operator)
):
    """Bulk upload subscribers from a CSV or XLSX file."""
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot upload subscribers")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")

    content = await file.read()
    filename = sanitize_filename(file.filename or "", default="subscribers").lower()
    rows = []

    try:
        if filename.endswith(".xlsx") or filename.endswith(".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb.active
            headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append({headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
        else:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            for row in reader:
                rows.append({k.strip().lower(): v.strip() for k, v in row.items()})
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse file: {e}")

    # Fetch operator plans for name→id mapping
    op_plans = await db.operator_plans.find(
        {"operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
    ).to_list(500)
    plan_map = {p["name"].strip().lower(): p for p in op_plans}

    # ── Pre-flight: check subscriber limit BEFORE processing any rows ──────────
    operator = await db.operators.find_one({"id": current_user["operator_id"], "deleted_at": None}, {"_id": 0})
    max_subscribers = None
    if operator and operator.get("saas_plan_id"):
        sp = await db.saas_plans.find_one({"id": operator["saas_plan_id"], "deleted_at": None}, {"_id": 0})
        if sp:
            max_subscribers = sp.get("max_subscribers")

    if max_subscribers is not None:
        current_count = await db.subscribers.count_documents(
            {"operator_id": current_user["operator_id"], "deleted_at": None}
        )
        # Count valid rows (name + whatsapp present) to get the intended upload size
        valid_row_count = sum(
            1 for r in rows
            if r.get("name", "").strip() and r.get("whatsapp_number", "").strip()
        )
        available_slots = max_subscribers - current_count
        if valid_row_count > available_slots:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Upload exceeds subscriber limit. "
                    f"Your plan allows {max_subscribers} subscribers. "
                    f"You currently have {current_count} and are trying to add {valid_row_count} more "
                    f"(total would be {current_count + valid_row_count}). "
                    f"Available slots: {available_slots}. "
                    f"Please upgrade your plan."
                )
            )

    now = datetime.now(timezone.utc)
    created, skipped, errors = [], [], []

    for idx, row in enumerate(rows, start=2):
        name = row.get("name", "").strip()
        whatsapp = row.get("whatsapp_number", "").strip()
        plan_name = row.get("plan_name", "").strip().lower()

        if not name or not whatsapp:
            errors.append({"row": idx, "reason": "name and whatsapp_number are required"})
            continue

        plan = plan_map.get(plan_name)
        if not plan:
            errors.append({"row": idx, "name": name, "reason": f"Plan '{row.get('plan_name','')}' not found"})
            continue

        # Check duplicate WhatsApp
        dup = await db.subscribers.find_one(
            {"whatsapp_number": whatsapp, "operator_id": current_user["operator_id"], "deleted_at": None}
        )
        if dup:
            skipped.append({"row": idx, "name": name, "reason": f"WhatsApp {whatsapp} already exists"})
            continue

        billing_date = int(row.get("billing_date", 1) or 1)
        billing_date = max(1, min(28, billing_date))
        discount = float(row.get("discount", 0) or 0)

        subscriber = {
            "id": generate_id(), "name": name, "whatsapp_number": whatsapp,
            "email": row.get("email", "") or None,
            "address": row.get("address", "") or None,
            "plan_id": plan["id"], "plan_name": plan["name"],
            "billing_date": billing_date, "discount": discount,
            "status": "active", "operator_id": current_user["operator_id"],
            "created_at": now.isoformat(), "updated_at": now.isoformat(), "deleted_at": None
        }
        await db.subscribers.insert_one(subscriber)
        created.append(name)

    return {
        "message": f"Bulk upload complete: {len(created)} created, {len(skipped)} skipped, {len(errors)} errors",
        "created": len(created), "skipped": len(skipped), "errors": errors[:20]
    }


# ─── Invoices ─────────────────────────────────────────────────────────────────

@router.post("/invoices", response_model=InvoiceResponse)
async def create_invoice(data: InvoiceCreate, request: Request, current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot create invoices")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    subscriber = await db.subscribers.find_one(
        {"id": data.subscriber_id, "operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
    )
    if not subscriber:
        raise HTTPException(status_code=404, detail="Subscriber not found")
    plan = await db.operator_plans.find_one({"id": data.plan_id, "deleted_at": None}, {"_id": 0})
    if not plan:
        raise HTTPException(status_code=404, detail="Plan not found")
    operator = await db.operators.find_one({"id": current_user["operator_id"], "deleted_at": None}, {"_id": 0})
    tax_amount = 0
    # GST can only be applied if operator has a valid GSTIN
    can_charge_gst = operator.get("charge_gst") and operator.get("gst_number")
    if can_charge_gst and plan.get("tax_percentage", 0) > 0:
        if plan.get("tax_type") == "exclusive":
            tax_amount = (data.base_amount - data.discount) * (plan["tax_percentage"] / 100)
        elif plan.get("tax_type") == "inclusive":
            tax_amount = (data.base_amount - data.discount) - ((data.base_amount - data.discount) / (1 + plan["tax_percentage"] / 100))
    final_amount = data.base_amount - data.discount + (tax_amount if plan.get("tax_type") == "exclusive" else 0)
    now = datetime.now(timezone.utc)
    invoice = {
        "id": generate_id(),
        "invoice_number": await generate_invoice_number_atomic(db),
        "subscriber_id": data.subscriber_id, "subscriber_name": subscriber["name"],
        "plan_id": data.plan_id, "plan_name": plan["name"],
        "base_amount": data.base_amount, "discount": data.discount,
        "tax_amount": round(tax_amount, 2), "final_amount": round(final_amount, 2),
        "service_start_date": data.service_start_date.isoformat(),
        "service_end_date": data.service_end_date.isoformat(),
        "due_date": data.due_date.isoformat(), "status": "pending", "payment_id": None,
        "operator_id": current_user["operator_id"],
        "created_at": now.isoformat(), "updated_at": now.isoformat(), "deleted_at": None
    }
    await db.invoices.insert_one(invoice)

    # Auto-send WhatsApp if whatsapp_notifications addon is active (uses platform WhatsApp config)
    auto_wa_sent = False
    # Build public invoice URL from request origin
    public_invoice_url = None
    try:
        # Get the origin from the request headers (set by browser)
        origin = request.headers.get("origin") or request.headers.get("referer", "")
        if origin:
            # Strip trailing slash and any path
            from urllib.parse import urlparse
            parsed = urlparse(origin)
            base_url = f"{parsed.scheme}://{parsed.netloc}"
        else:
            # Fallback: use request base URL (strips /api prefix)
            base_url = str(request.base_url).rstrip("/")
        if base_url:
            public_invoice_url = f"{base_url}/invoice/{invoice['id']}"
    except Exception:
        pass

    if await _has_addon(current_user["operator_id"], "whatsapp_notifications"):
        try:
            wa_config = await _get_platform_whatsapp_config()
            if wa_config:
                template_settings = await _get_whatsapp_template_settings()
                template_name = template_settings.get("invoice_template") or "invoice_notification"
                from services.whatsapp_service import WhatsAppService
                wa_service = WhatsAppService(wa_config["phone_number_id"], wa_config["access_token"])
                await wa_service.send_invoice_notification(
                    recipient_phone=subscriber["whatsapp_number"],
                    customer_name=subscriber["name"],
                    invoice_number=invoice["invoice_number"],
                    amount=f"₹{invoice['final_amount']:,.2f}",
                    due_date=data.due_date.strftime("%d %b %Y"),
                    payment_link=public_invoice_url,
                    template_name_override=template_name
                )
                auto_wa_sent = True
        except Exception as e:
            logger.warning(f"Auto WhatsApp send failed: {e}")

    response = InvoiceResponse(**{
        **invoice, "created_at": now,
        "service_start_date": data.service_start_date,
        "service_end_date": data.service_end_date,
        "due_date": data.due_date
    })
    # Return extra meta for frontend to decide WhatsApp Web button visibility
    result = response.model_dump()
    result["auto_wa_sent"] = auto_wa_sent
    result["has_whatsapp_addon"] = await _has_addon(current_user["operator_id"], "whatsapp_notifications")
    result["public_url"] = public_invoice_url
    return result


@router.get("/invoices", response_model=List[InvoiceResponse])
async def get_invoices(
    status: Optional[str] = None, subscriber_id: Optional[str] = None,
    current_user: dict = Depends(require_operator)
):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot access operator invoices")
    query = {"operator_id": current_user["operator_id"], "deleted_at": None}
    if status:
        query["status"] = status
    if subscriber_id:
        query["subscriber_id"] = subscriber_id
    invoices = await db.invoices.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return [
        InvoiceResponse(**{
            **inv, "created_at": datetime.fromisoformat(inv["created_at"]),
            "service_start_date": datetime.fromisoformat(inv["service_start_date"]),
            "service_end_date": datetime.fromisoformat(inv["service_end_date"]),
            "due_date": datetime.fromisoformat(inv["due_date"])
        })
        for inv in invoices
    ]


@router.put("/invoices/{invoice_id}/status")
async def update_invoice_status(
    invoice_id: str, status: str = Query(...), current_user: dict = Depends(require_operator)
):
    status = sanitize_text(status)
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    if status not in ["pending", "paid", "overdue", "cancelled"]:
        raise HTTPException(status_code=400, detail="Invalid status")
    result = await db.invoices.update_one(
        {"id": invoice_id, "operator_id": current_user["operator_id"], "deleted_at": None},
        {"$set": {"status": status, "updated_at": datetime.now(timezone.utc).isoformat()}}
    )
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Invoice not found")
    return {"message": f"Invoice marked as {status}"}


@router.post("/invoices/{invoice_id}/payment-link", response_model=PaymentLinkResponse)
async def create_payment_link(invoice_id: str, current_user: dict = Depends(require_operator)):
    from services.razorpay_service import RazorpayService
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    invoice = await db.invoices.find_one(
        {"id": invoice_id, "operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
    )
    if not invoice:
        raise HTTPException(status_code=404, detail="Invoice not found")

    # Determine which payment gateway to use:
    # - custom_payment_gateway addon → operator's own keys
    # - payment_gateway addon → platform keys
    has_custom_pg = await _has_addon(current_user["operator_id"], "custom_payment_gateway")
    has_platform_pg = await _has_addon(current_user["operator_id"], "payment_gateway")

    if has_custom_pg:
        # Use operator's own gateway keys
        gateway = await db.payment_gateways.find_one({"operator_id": current_user["operator_id"]}, {"_id": 0})
        if not gateway or not gateway.get("is_active"):
            raise HTTPException(status_code=400, detail="Custom payment gateway not configured. Please set up your gateway keys in Settings.")
    elif has_platform_pg:
        # Use platform gateway keys
        gateway = await db.payment_gateways.find_one({"is_platform_gateway": True, "is_active": True}, {"_id": 0})
        if not gateway:
            raise HTTPException(status_code=400, detail="Platform payment gateway not configured. Please contact admin.")
    else:
        raise HTTPException(status_code=403, detail="Payment gateway add-on is not enabled. Please activate 'Payment Gateway' or 'Custom Payment Gateway' add-on.")

    subscriber = await db.subscribers.find_one({"id": invoice["subscriber_id"], "deleted_at": None}, {"_id": 0})
    try:
        razorpay_service = RazorpayService(gateway["api_key"], gateway["api_secret"])
        payment_link = razorpay_service.create_payment_link(
            amount=invoice["final_amount"],
            description=f"Invoice {invoice['invoice_number']}",
            customer_name=subscriber["name"] if subscriber else "",
            customer_email=subscriber.get("email", "") if subscriber else "",
            customer_phone=subscriber.get("whatsapp_number", "") if subscriber else "",
            invoice_number=invoice["invoice_number"]
        )
        qr_code = razorpay_service.generate_qr_code(payment_link["short_url"])
        await db.invoices.update_one(
            {"id": invoice_id},
            {"$set": {"payment_link": payment_link["short_url"], "payment_link_id": payment_link["id"],
                      "updated_at": datetime.now(timezone.utc).isoformat()}}
        )
        return PaymentLinkResponse(
            payment_link=payment_link["short_url"], payment_link_id=payment_link["id"],
            qr_code=qr_code, amount=invoice["final_amount"]
        )
    except Exception as e:
        logger.error(f"Payment link creation failed: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to create payment link: {e}")


@router.get("/invoices/{invoice_id}/pdf")
async def get_invoice_pdf(invoice_id: str, current_user: dict = Depends(require_operator)):
    from services.pdf_service import InvoicePDFService
    from services.razorpay_service import RazorpayService
    invoice = await db.invoices.find_one(
        {"id": invoice_id, "operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
    )
    if not invoice:
        raise HTTPException(status_code=404, detail="Invoice not found")
    operator = await db.operators.find_one({"id": current_user["operator_id"], "deleted_at": None}, {"_id": 0})
    subscriber = await db.subscribers.find_one({"id": invoice["subscriber_id"], "deleted_at": None}, {"_id": 0})
    plan = await db.operator_plans.find_one({"id": invoice["plan_id"], "deleted_at": None}, {"_id": 0})
    qr_code = None
    if invoice.get("payment_link"):
        gateway = await db.payment_gateways.find_one({"operator_id": current_user["operator_id"]}, {"_id": 0})
        if gateway:
            try:
                razorpay_service = RazorpayService(gateway["api_key"], gateway["api_secret"])
                qr_code = razorpay_service.generate_qr_code(invoice["payment_link"])
            except Exception:
                pass
    pdf_service = InvoicePDFService()
    inv_settings = await db.invoice_settings.find_one(
        {"operator_id": current_user["operator_id"]}, {"_id": 0}
    )
    template = (inv_settings or {}).get("invoice_template", "classic")
    pdf_bytes = pdf_service.generate_invoice_pdf(
        invoice_data=invoice, operator_data={**(operator or {}), **(inv_settings or {})},
        subscriber_data=subscriber or {}, plan_data=plan or {}, qr_code_base64=qr_code,
        template=template,
    )
    return Response(
        content=pdf_bytes, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=Invoice_{invoice['invoice_number']}.pdf"}
    )


# ─── Staff ────────────────────────────────────────────────────────────────────

@router.post("/staff", response_model=StaffResponse)
async def create_staff(data: StaffCreate, current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot create staff")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    operator = await db.operators.find_one({"id": current_user["operator_id"], "deleted_at": None}, {"_id": 0})
    if operator:
        # Block staff creation on trial plan with a clear message
        if operator.get("status") == "trial":
            raise HTTPException(status_code=403, detail="Please subscribe to use this feature.")
        # Use operator-level max_staff override (set when staff_management addon is assigned)
        # falling back to the SaaS plan's max_staff
        operator_max_staff = operator.get("max_staff")
        if operator_max_staff is None:
            plan = await db.saas_plans.find_one({"id": operator.get("saas_plan_id"), "deleted_at": None}, {"_id": 0})
            operator_max_staff = plan["max_staff"] if plan else 0
        current_count = await db.users.count_documents(
            {"operator_id": current_user["operator_id"], "role": "staff", "deleted_at": None}
        )
        if current_count >= operator_max_staff:
            raise HTTPException(status_code=403, detail=f"Staff limit ({operator_max_staff}) reached")
    existing = await db.users.find_one({"email": data.email, "deleted_at": None})
    if existing:
        raise HTTPException(status_code=400, detail="Email already registered")
    now = datetime.now(timezone.utc)
    user_id = generate_id()
    user = {
        "id": user_id, "email": data.email, "name": data.name, "phone": data.phone,
        "password": hash_password(data.password), "role": "staff",
        "permissions": data.permissions, "operator_id": current_user["operator_id"],
        "status": "active", "created_at": now.isoformat(),
        "updated_at": now.isoformat(), "deleted_at": None
    }
    await db.users.insert_one(user)
    return StaffResponse(
        id=user_id, name=data.name, email=data.email, phone=data.phone,
        role="staff", permissions=data.permissions,
        operator_id=current_user["operator_id"], status="active", created_at=now
    )


@router.get("/staff", response_model=List[StaffResponse])
async def get_staff(current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot access operator staff")
    staff = await db.users.find(
        {"operator_id": current_user["operator_id"], "role": "staff", "deleted_at": None},
        {"_id": 0, "password": 0}
    ).to_list(100)
    return [StaffResponse(**{**s, "permissions": s.get("permissions", []),
                             "created_at": datetime.fromisoformat(s["created_at"])}) for s in staff]


@router.delete("/staff/{staff_id}")
async def delete_staff(staff_id: str, current_user: dict = Depends(require_operator_no_staff)):
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    result = await db.users.update_one(
        {"id": staff_id, "operator_id": current_user["operator_id"], "role": "staff", "deleted_at": None},
        {"$set": {"deleted_at": datetime.now(timezone.utc).isoformat()}}
    )
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Staff not found")
    return {"message": "Staff deleted"}


# ─── Payment Gateway ────────────────────────────────────────────────────────

@router.post("/payment-gateway")
async def configure_payment_gateway(data: PaymentGatewayConfig, current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot configure operator payment gateway")
    if await check_operator_read_only(current_user["operator_id"]):
        raise HTTPException(status_code=403, detail="Account is in read-only mode")
    if not await _has_addon(current_user["operator_id"], "custom_payment_gateway"):
        raise HTTPException(status_code=403, detail="Custom Payment Gateway add-on is not enabled for your plan.")
    now = datetime.now(timezone.utc)
    gateway_config = {
        "id": generate_id(), "operator_id": current_user["operator_id"],
        "gateway_type": data.gateway_type, "api_key": data.api_key,
        "api_secret": data.api_secret, "webhook_secret": data.webhook_secret,
        "is_active": True, "created_at": now.isoformat(), "updated_at": now.isoformat()
    }
    await db.payment_gateways.update_one(
        {"operator_id": current_user["operator_id"]}, {"$set": gateway_config}, upsert=True
    )
    return {"message": "Payment gateway configured successfully"}


@router.get("/payment-gateway")
async def get_payment_gateway(current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot access operator payment gateway")
    gateway = await db.payment_gateways.find_one(
        {"operator_id": current_user["operator_id"]}, {"_id": 0, "api_secret": 0}
    )
    if not gateway:
        return {"configured": False}
    return {
        "configured": True, "gateway_type": gateway["gateway_type"],
        "api_key": gateway["api_key"][:8] + "****", "is_active": gateway.get("is_active", False)
    }


# ─── Reports ─────────────────────────────────────────────────────────────────

@router.get("/reports/revenue")
async def get_revenue_report(
    start_date: Optional[str] = None, end_date: Optional[str] = None,
    current_user: dict = Depends(require_operator)
):
    start_date = sanitize_text(start_date) if start_date else start_date
    end_date = sanitize_text(end_date) if end_date else end_date
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot access operator reports")
    query = {"operator_id": current_user["operator_id"], "status": "paid", "deleted_at": None}
    if start_date:
        query["created_at"] = {"$gte": start_date}
    if end_date:
        query.setdefault("created_at", {})["$lte"] = end_date
    invoices = await db.invoices.find(query, {"_id": 0}).to_list(10000)
    return {
        "total_invoices": len(invoices),
        "total_revenue": round(sum(inv.get("final_amount", 0) for inv in invoices), 2),
        "total_base_amount": round(sum(inv.get("base_amount", 0) for inv in invoices), 2),
        "total_tax": round(sum(inv.get("tax_amount", 0) for inv in invoices), 2),
        "total_discount": round(sum(inv.get("discount", 0) for inv in invoices), 2),
    }


@router.get("/reports/gst-summary")
async def get_gst_summary(
    start_date: Optional[str] = None, end_date: Optional[str] = None,
    current_user: dict = Depends(require_operator)
):
    start_date = sanitize_text(start_date) if start_date else start_date
    end_date = sanitize_text(end_date) if end_date else end_date
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot access operator reports")
    query = {"operator_id": current_user["operator_id"], "deleted_at": None}
    if start_date:
        query["created_at"] = {"$gte": start_date}
    if end_date:
        query.setdefault("created_at", {})["$lte"] = end_date
    invoices = await db.invoices.find(query, {"_id": 0}).to_list(10000)
    total_taxable = sum(inv.get("base_amount", 0) - inv.get("discount", 0) for inv in invoices)
    total_gst = sum(inv.get("tax_amount", 0) for inv in invoices)
    return {
        "total_taxable_amount": round(total_taxable, 2), "total_gst_collected": round(total_gst, 2),
        "cgst": round(total_gst / 2, 2), "sgst": round(total_gst / 2, 2)
    }


@router.get("/reports/pending-overdue")
async def get_pending_overdue_report(current_user: dict = Depends(require_operator)):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot access operator reports")
    pending = await db.invoices.find(
        {"operator_id": current_user["operator_id"], "status": "pending", "deleted_at": None}, {"_id": 0}
    ).to_list(1000)
    overdue = await db.invoices.find(
        {"operator_id": current_user["operator_id"], "status": "overdue", "deleted_at": None}, {"_id": 0}
    ).to_list(1000)
    return {
        "pending_count": len(pending), "pending_amount": round(sum(inv.get("final_amount", 0) for inv in pending), 2),
        "overdue_count": len(overdue), "overdue_amount": round(sum(inv.get("final_amount", 0) for inv in overdue), 2)
    }


@router.get("/reports/invoices")
async def get_report_invoices(
    page: int = Query(1, ge=1),
    limit: int = Query(15, ge=1, le=100),
    status: Optional[str] = None,
    search: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    sort_by: Optional[str] = "created_at",
    sort_order: Optional[str] = "desc",
    current_user: dict = Depends(require_operator),
):
    """Get paginated invoice list with filters for reports."""
    status = sanitize_text(status) if status else status
    search = sanitize_text(search) if search else search
    start_date = sanitize_text(start_date) if start_date else start_date
    end_date = sanitize_text(end_date) if end_date else end_date
    sort_by = sanitize_text(sort_by) if sort_by else sort_by
    sort_order = sanitize_text(sort_order) if sort_order else sort_order
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Admin cannot access operator reports")
    query = {"operator_id": current_user["operator_id"], "deleted_at": None}

    # Status filter
    if status and status != "all":
        query["status"] = status

    # Date range filter on created_at
    if start_date or end_date:
        date_filter = {}
        if start_date:
            date_filter["$gte"] = start_date
        if end_date:
            date_filter["$lte"] = end_date + "T23:59:59"
        query["created_at"] = date_filter

    # Text search on invoice_number or subscriber_name
    if search and search.strip():
        search_regex = {"$regex": search.strip(), "$options": "i"}
        query["$or"] = [
            {"invoice_number": search_regex},
            {"subscriber_name": search_regex},
        ]

    # Sort
    sort_dir = -1 if sort_order == "desc" else 1
    valid_sort_fields = {"created_at", "final_amount", "due_date", "invoice_number", "status"}
    sort_field = sort_by if sort_by in valid_sort_fields else "created_at"

    # Count total
    total = await db.invoices.count_documents(query)

    # Paginate
    skip = (page - 1) * limit
    invoices = await db.invoices.find(query, {"_id": 0}).sort(
        sort_field, sort_dir
    ).skip(skip).limit(limit).to_list(limit)

    return {
        "invoices": invoices,
        "total": total,
        "page": page,
        "limit": limit,
        "total_pages": max(1, -(-total // limit)),  # ceil division
    }


# ─── Audit Logs ────────────────────────────────────────────────────────────

@router.get("/audit-logs", response_model=List[AuditLogResponse])
async def get_operator_audit_logs(
    skip: int = 0, limit: int = 50, current_user: dict = Depends(require_operator)
):
    if current_user["role"] == "admin":
        raise HTTPException(status_code=400, detail="Use admin audit logs endpoint")
    if not await _has_addon(current_user["operator_id"], "audit_log"):
        raise HTTPException(status_code=403, detail="Audit Logs add-on is not enabled for your plan.")
    logs = await db.audit_logs.find(
        {"operator_id": current_user["operator_id"]}, {"_id": 0}
    ).sort("created_at", -1).skip(skip).limit(limit).to_list(limit)
    return [AuditLogResponse(**{**log, "created_at": datetime.fromisoformat(log["created_at"])}) for log in logs]


# ─── Operator Settlements ───────────────────────────────────────────────────────

@router.get("/settlements/summary")
async def get_operator_settlements_summary(current_user: dict = Depends(require_operator)):
    """Get settlement summary for operator."""
    now = datetime.now(timezone.utc)
    start_of_month = now.replace(day=1).strftime("%Y-%m-%d")

    all_settlements = await db.settlements.find(
        {"operator_id": current_user["operator_id"]}, {"_id": 0}
    ).to_list(10000)

    total_settled = sum(s["net_settlement"] for s in all_settlements if s["status"] == "completed")
    total_pending = sum(s["net_settlement"] for s in all_settlements if s["status"] == "pending")
    total_platform_fee = sum(s["platform_fee"] for s in all_settlements if s["status"] == "completed")
    total_collections = sum(s["total_collections"] for s in all_settlements if s["status"] == "completed")

    month_settlements = [s for s in all_settlements if s["settlement_date"] >= start_of_month]
    month_settled = sum(s["net_settlement"] for s in month_settlements if s["status"] == "completed")
    month_pending = sum(s["net_settlement"] for s in month_settlements if s["status"] == "pending")

    pending_count = sum(1 for s in all_settlements if s["status"] == "pending")
    completed_count = sum(1 for s in all_settlements if s["status"] == "completed")

    return {
        "total_settled": total_settled,
        "total_pending": total_pending,
        "total_platform_fee": total_platform_fee,
        "total_collections": total_collections,
        "month_settled": month_settled,
        "month_pending": month_pending,
        "pending_count": pending_count,
        "completed_count": completed_count,
        "total_count": len(all_settlements),
    }


@router.get("/settlements")
async def get_operator_settlements(
    status: Optional[str] = None,
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
    current_user: dict = Depends(require_operator),
):
    """List settlements for operator."""
    query = {"operator_id": current_user["operator_id"]}
    if status:
        query["status"] = status

    total = await db.settlements.count_documents(query)
    skip = (page - 1) * limit

    settlements = await db.settlements.find(query, {"_id": 0}).sort(
        "settlement_date", -1
    ).skip(skip).limit(limit).to_list(limit)

    return {
        "settlements": settlements,
        "total": total,
        "page": page,
        "pages": (total + limit - 1) // limit,
    }


@router.get("/settlements/{settlement_id}")
async def get_operator_settlement_detail(settlement_id: str, current_user: dict = Depends(require_operator)):
    """Get settlement detail with included invoices."""
    settlement = await db.settlements.find_one(
        {"id": settlement_id, "operator_id": current_user["operator_id"]}, {"_id": 0}
    )
    if not settlement:
        raise HTTPException(status_code=404, detail="Settlement not found")

    # Fetch the invoices in this settlement
    invoices = await db.invoices.find(
        {"id": {"$in": settlement.get("invoice_ids", [])}},
        {"_id": 0, "id": 1, "invoice_number": 1, "subscriber_name": 1,
         "final_amount": 1, "tax_amount": 1, "base_amount": 1, "discount": 1,
         "paid_at": 1, "plan_name": 1}
    ).to_list(1000)

    # Fetch operator bank details
    operator = await db.operators.find_one(
        {"id": current_user["operator_id"], "deleted_at": None},
        {"_id": 0, "bank_account_name": 1, "bank_account_number": 1,
         "bank_ifsc": 1, "bank_name": 1, "company_name": 1}
    )

    return {
        **settlement,
        "invoices": invoices,
        "operator_details": operator,
    }


# ─── WhatsApp Notifications (uses platform global WhatsApp config) ────────────

@router.post("/send-notification")
async def send_whatsapp_notification(data: SendNotificationRequest, current_user: dict = Depends(require_operator)):
    from services.whatsapp_service import WhatsAppService
    wa_config = await _get_platform_whatsapp_config()
    if not wa_config:
        raise HTTPException(status_code=400, detail="WhatsApp not configured. Please contact admin.")
    invoice = await db.invoices.find_one(
        {"id": data.invoice_id, "operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
    )
    if not invoice:
        raise HTTPException(status_code=404, detail="Invoice not found")
    subscriber = await db.subscribers.find_one({"id": invoice["subscriber_id"], "deleted_at": None}, {"_id": 0})
    if not subscriber:
        raise HTTPException(status_code=404, detail="Subscriber not found")
    try:
        template_settings = await _get_whatsapp_template_settings()
        wa_service = WhatsAppService(wa_config["phone_number_id"], wa_config["access_token"])
        if data.notification_type == "reminder":
            template_name = template_settings.get("reminder_template") or "payment_reminder"
            due_date = datetime.fromisoformat(invoice["due_date"].replace('Z', '+00:00'))
            days_overdue = max(0, (datetime.now(timezone.utc) - due_date).days)
            result = await wa_service.send_payment_reminder(
                recipient_phone=subscriber["whatsapp_number"], customer_name=subscriber["name"],
                invoice_number=invoice["invoice_number"],
                amount_due=f"₹{invoice['final_amount']:,.2f}", days_overdue=str(days_overdue),
                payment_link=invoice.get("payment_link"),
                template_name_override=template_name
            )
        else:
            template_name = template_settings.get("invoice_template") or "invoice_notification"
            result = await wa_service.send_invoice_notification(
                recipient_phone=subscriber["whatsapp_number"], customer_name=subscriber["name"],
                invoice_number=invoice["invoice_number"],
                amount=f"₹{invoice['final_amount']:,.2f}",
                due_date=datetime.fromisoformat(invoice["due_date"].replace('Z', '+00:00')).strftime("%d %b %Y"),
                payment_link=invoice.get("payment_link"),
                template_name_override=template_name
            )
        return {"success": True, "message_id": result.get("messages", [{}])[0].get("id")}
    except Exception as e:
        logger.error(f"WhatsApp notification failed: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to send notification: {e}")


@router.post("/bulk-notification")
async def send_bulk_notification(data: BulkNotificationRequest, current_user: dict = Depends(require_operator)):
    from services.whatsapp_service import WhatsAppService
    wa_config = await _get_platform_whatsapp_config()
    if not wa_config:
        raise HTTPException(status_code=400, detail="WhatsApp not configured. Please contact admin.")
    template_settings = await _get_whatsapp_template_settings()
    invoice_template = template_settings.get("invoice_template") or "invoice_notification"
    results = {"sent": 0, "failed": 0, "errors": []}
    wa_service = WhatsAppService(wa_config["phone_number_id"], wa_config["access_token"])
    for subscriber_id in data.subscriber_ids:
        try:
            subscriber = await db.subscribers.find_one(
                {"id": subscriber_id, "operator_id": current_user["operator_id"], "deleted_at": None}, {"_id": 0}
            )
            if not subscriber:
                continue
            invoice = await db.invoices.find_one(
                {"subscriber_id": subscriber_id, "status": {"$in": ["pending", "overdue"]}, "deleted_at": None},
                {"_id": 0}
            )
            if invoice:
                await wa_service.send_invoice_notification(
                    recipient_phone=subscriber["whatsapp_number"], customer_name=subscriber["name"],
                    invoice_number=invoice["invoice_number"],
                    amount=f"₹{invoice['final_amount']:,.2f}",
                    due_date=datetime.fromisoformat(invoice["due_date"].replace('Z', '+00:00')).strftime("%d %b %Y"),
                    payment_link=invoice.get("payment_link"),
                    template_name_override=invoice_template
                )
                results["sent"] += 1
        except Exception as e:
            results["failed"] += 1
            results["errors"].append({"subscriber_id": subscriber_id, "error": str(e)})
    return results


# ─── Reminder Settings ──────────────────────────────────────────────────────

VALID_BEFORE_DAYS = [1, 2, 3, 5, 7]
VALID_AFTER_DAYS = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10]


@router.get("/reminder-settings")
async def get_reminder_settings(current_user: dict = Depends(require_operator)):
    """Get operator's payment reminder automation settings."""
    operator_id = current_user["operator_id"]
    if not await _has_addon(operator_id, "whatsapp_notifications"):
        raise HTTPException(status_code=403, detail="WhatsApp notifications add-on is not enabled")

    doc = await db.reminder_settings.find_one({"operator_id": operator_id}, {"_id": 0})
    if not doc:
        # Auto-create default schedule: 7,5,3,2,1 days before + on due + 1-10 days after
        now = datetime.now(timezone.utc).isoformat()
        default_settings = {
            "operator_id": operator_id,
            "enabled": True,
            "remind_before_due": [7, 5, 3, 2, 1],
            "remind_on_due": True,
            "remind_after_due": [1, 2, 3, 4, 5, 6, 7, 8, 9, 10],
            "max_reminders_per_invoice": 20,
            "created_at": now,
            "updated_at": now,
        }
        await db.reminder_settings.insert_one({**default_settings})
        return default_settings
    doc.pop("_id", None)
    return doc


@router.put("/reminder-settings")
async def update_reminder_settings(
    data: ReminderSettingsUpdate,
    current_user: dict = Depends(require_operator),
):
    """Update operator's payment reminder automation settings."""
    operator_id = current_user["operator_id"]
    if not await _has_addon(operator_id, "whatsapp_notifications"):
        raise HTTPException(status_code=403, detail="WhatsApp notifications add-on is not enabled")
    await check_operator_read_only(current_user)

    # Validate day values
    for d in data.remind_before_due:
        if d not in VALID_BEFORE_DAYS:
            raise HTTPException(status_code=400, detail=f"Invalid remind_before_due day: {d}. Allowed: {VALID_BEFORE_DAYS}")
    for d in data.remind_after_due:
        if d not in VALID_AFTER_DAYS:
            raise HTTPException(status_code=400, detail=f"Invalid remind_after_due day: {d}. Allowed: {VALID_AFTER_DAYS}")

    now = datetime.now(timezone.utc).isoformat()
    update_doc = {
        "operator_id": operator_id,
        "enabled": data.enabled,
        "remind_before_due": sorted(set(data.remind_before_due), reverse=True),
        "remind_on_due": data.remind_on_due,
        "remind_after_due": sorted(set(data.remind_after_due)),
        "max_reminders_per_invoice": max(1, min(data.max_reminders_per_invoice, 20)),
        "updated_at": now,
    }

    await db.reminder_settings.update_one(
        {"operator_id": operator_id},
        {"$set": update_doc, "$setOnInsert": {"created_at": now}},
        upsert=True,
    )

    await log_audit(
        current_user["id"], current_user["name"], current_user["role"],
        "update", "reminder_settings", None, update_doc,
        operator_id=operator_id,
    )

    return {**update_doc, "message": "Reminder settings updated successfully"}
