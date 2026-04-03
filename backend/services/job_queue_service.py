"""Background job queue service using MongoDB as the job store.

Supports bulk uploads and bulk notifications with status tracking.
Jobs are processed by APScheduler every 30 seconds.
"""

import logging
from datetime import datetime, timezone
from typing import Optional, Any
import io
import csv

from database import db
from utils import generate_id, generate_invoice_number_atomic
from sanitization import sanitize_filename
from audit import log_audit

logger = logging.getLogger(__name__)


class JobQueueService:
    """Service for managing background jobs."""

    @staticmethod
    async def enqueue_job(job_type: str, operator_id: str, user_id: str, payload: dict) -> str:
        """
        Submit a job to the queue.

        Args:
            job_type: Type of job (bulk_upload_plans, bulk_upload_subscribers, bulk_upload_invoices, bulk_notification)
            operator_id: ID of the operator submitting the job
            user_id: ID of the user submitting the job
            payload: Job-specific data (file_content, subscriber_ids, etc.)

        Returns:
            job_id: Unique job identifier
        """
        job_id = generate_id()
        job = {
            "id": job_id,
            "type": job_type,
            "status": "pending",
            "operator_id": operator_id,
            "created_by": user_id,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "started_at": None,
            "completed_at": None,
            "payload": payload,
            "result": None,
            "error": None,
        }
        await db.background_jobs.insert_one(job)
        logger.info(f"Enqueued job {job_id} of type {job_type} for operator {operator_id}")
        return job_id

    @staticmethod
    async def get_job(job_id: str, operator_id: str) -> Optional[dict]:
        """
        Retrieve job status by ID (scoped to operator for security).

        Args:
            job_id: Job identifier
            operator_id: Operator ID (for access control)

        Returns:
            Job document or None if not found
        """
        job = await db.background_jobs.find_one(
            {"id": job_id, "operator_id": operator_id},
            {"_id": 0}
        )
        return job

    @staticmethod
    async def process_next_job():
        """
        Process one pending job from the queue.
        Called by APScheduler every 30 seconds.
        """
        try:
            # Claim one pending job atomically
            job = await db.background_jobs.find_one_and_update(
                {"status": "pending"},
                {
                    "$set": {
                        "status": "processing",
                        "started_at": datetime.now(timezone.utc).isoformat(),
                    }
                },
                {"_id": 0},
            )

            if not job:
                return  # No pending jobs

            job_id = job["id"]
            job_type = job["type"]
            operator_id = job["operator_id"]

            logger.info(f"Processing job {job_id} (type={job_type}, operator={operator_id})")

            try:
                # Dispatch to appropriate handler
                if job_type == "bulk_upload_plans":
                    result = await JobQueueService._run_bulk_plans(job)
                elif job_type == "bulk_upload_subscribers":
                    result = await JobQueueService._run_bulk_subscribers(job)
                elif job_type == "bulk_upload_invoices":
                    result = await JobQueueService._run_bulk_invoices(job)
                elif job_type == "bulk_notification":
                    result = await JobQueueService._run_bulk_notification(job)
                else:
                    raise ValueError(f"Unknown job type: {job_type}")

                # Mark as completed
                await db.background_jobs.update_one(
                    {"id": job_id},
                    {
                        "$set": {
                            "status": "completed",
                            "completed_at": datetime.now(timezone.utc).isoformat(),
                            "result": result,
                        }
                    },
                )
                logger.info(f"Job {job_id} completed successfully")

                # Log to audit logs
                user = await db.users.find_one({"id": job["created_by"]}, {"_id": 0, "name": 1})
                await log_audit(
                    user_id=job["created_by"],
                    user_name=user.get("name", "Unknown") if user else "Unknown",
                    role="operator",
                    action="bulk_upload_completed",
                    module=job_type.replace("bulk_upload_", "").replace("bulk_", "").capitalize(),
                    new_value=result,
                    operator_id=operator_id
                )

            except Exception as e:
                # Mark as failed
                error_msg = str(e)
                logger.error(f"Job {job_id} failed: {error_msg}", exc_info=True)
                await db.background_jobs.update_one(
                    {"id": job_id},
                    {
                        "$set": {
                            "status": "failed",
                            "completed_at": datetime.now(timezone.utc).isoformat(),
                            "error": error_msg,
                        }
                    },
                )

                # Log to audit logs
                user = await db.users.find_one({"id": job["created_by"]}, {"_id": 0, "name": 1})
                await log_audit(
                    user_id=job["created_by"],
                    user_name=user.get("name", "Unknown") if user else "Unknown",
                    role="operator",
                    action="bulk_upload_failed",
                    module=job_type.replace("bulk_upload_", "").replace("bulk_", "").capitalize(),
                    new_value={"error": error_msg},
                    operator_id=operator_id
                )

        except Exception as e:
            logger.error(f"Unexpected error in job processor: {e}", exc_info=True)

    # ─── Job Handlers ────────────────────────────────────────────────────────

    @staticmethod
    async def _run_bulk_plans(job: dict) -> dict:
        """Handler for bulk_upload_plans."""
        payload = job["payload"]
        operator_id = job["operator_id"]
        file_content = payload.get("file_content")
        filename = payload.get("filename", "plans")

        if isinstance(file_content, str):
            file_content = file_content.encode("utf-8")

        rows = []
        VALID_VALIDITY = ["monthly", "quarterly", "half_yearly", "yearly"]
        VALID_TAX_TYPE = ["inclusive", "exclusive", "none"]

        # Parse file
        if filename.lower().endswith((".xlsx", ".xls")):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_content))
            ws = wb.active
            headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append({headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
        else:
            reader = csv.DictReader(io.StringIO(file_content.decode("utf-8-sig")))
            for row in reader:
                rows.append({k.strip().lower(): v.strip() for k, v in row.items()})

        now = datetime.now(timezone.utc)
        created, skipped, errors = [], [], []

        for idx, row in enumerate(rows, start=2):
            name = row.get("name", "").strip()
            if not name:
                errors.append({"row": idx, "reason": "name is required"})
                continue

            try:
                price = float(row.get("price", 0) or 0)
            except ValueError:
                errors.append({"row": idx, "name": name, "reason": "Invalid price"})
                continue

            validity = row.get("validity", "monthly").strip().lower()
            if validity not in VALID_VALIDITY:
                errors.append({"row": idx, "name": name, "reason": f"Invalid validity '{validity}'. Use: {VALID_VALIDITY}"})
                continue

            tax_type = row.get("tax_type", "none").strip().lower()
            if tax_type not in VALID_TAX_TYPE:
                tax_type = "none"

            try:
                tax_percentage = float(row.get("tax_percentage", 0) or 0)
            except ValueError:
                tax_percentage = 0

            # Check for duplicate name
            dup = await db.operator_plans.find_one(
                {"name": name, "operator_id": operator_id, "deleted_at": None}
            )
            if dup:
                skipped.append({"row": idx, "name": name, "reason": "Plan name already exists"})
                continue

            plan = {
                "id": generate_id(),
                "name": name,
                "price": price,
                "validity": validity,
                "tax_percentage": tax_percentage,
                "tax_type": tax_type,
                "description": row.get("description", "") or None,
                "status": "active",
                "operator_id": operator_id,
                "created_at": now.isoformat(),
                "updated_at": now.isoformat(),
                "deleted_at": None,
            }
            await db.operator_plans.insert_one(plan)
            created.append(name)

        return {
            "message": f"Bulk upload complete: {len(created)} created, {len(skipped)} skipped, {len(errors)} errors",
            "created": len(created),
            "skipped": len(skipped),
            "errors": errors[:20],
        }

    @staticmethod
    async def _run_bulk_subscribers(job: dict) -> dict:
        """Handler for bulk_upload_subscribers."""
        payload = job["payload"]
        operator_id = job["operator_id"]
        file_content = payload.get("file_content")
        filename = payload.get("filename", "subscribers")

        if isinstance(file_content, str):
            file_content = file_content.encode("utf-8")

        rows = []

        # Parse file
        if filename.lower().endswith((".xlsx", ".xls")):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_content))
            ws = wb.active
            headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append({headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
        else:
            reader = csv.DictReader(io.StringIO(file_content.decode("utf-8-sig")))
            for row in reader:
                rows.append({k.strip().lower(): v.strip() for k, v in row.items()})

        # Fetch operator plans for name→id mapping
        op_plans = await db.operator_plans.find(
            {"operator_id": operator_id, "deleted_at": None}, {"_id": 0}
        ).to_list(500)
        plan_map = {p["name"].strip().lower(): p for p in op_plans}

        # Pre-flight: check subscriber limit
        operator = await db.operators.find_one({"id": operator_id, "deleted_at": None}, {"_id": 0})
        max_subscribers = None
        if operator and operator.get("saas_plan_id"):
            sp = await db.saas_plans.find_one({"id": operator["saas_plan_id"], "deleted_at": None}, {"_id": 0})
            if sp:
                max_subscribers = sp.get("max_subscribers")

        if max_subscribers is not None:
            current_count = await db.subscribers.count_documents(
                {"operator_id": operator_id, "deleted_at": None}
            )
            valid_row_count = sum(
                1 for r in rows
                if r.get("name", "").strip() and r.get("whatsapp_number", "").strip()
            )
            available_slots = max_subscribers - current_count
            if valid_row_count > available_slots:
                raise ValueError(
                    f"Upload exceeds subscriber limit. Your plan allows {max_subscribers} subscribers. "
                    f"You currently have {current_count} and are trying to add {valid_row_count} more "
                    f"(total would be {current_count + valid_row_count}). Available slots: {available_slots}. "
                    f"Please upgrade your plan."
                )

        now = datetime.now(timezone.utc)
        created, skipped, errors = [], [], []

        for idx, row in enumerate(rows, start=2):
            name = row.get("name", "").strip()
            whatsapp = row.get("whatsapp_number", "").strip()

            if not name or not whatsapp:
                errors.append({"row": idx, "reason": "name and whatsapp_number are required"})
                continue

            # Parse up to 5 plan slots (plan_name_1..5)
            plans = []
            plan_errors = []
            for i in range(1, 6):
                pname = row.get(f"plan_name_{i}", "").strip()
                if not pname:
                    continue
                plan = plan_map.get(pname.lower())
                if not plan:
                    plan_errors.append(f"Plan '{pname}' not found")
                    continue
                try:
                    bdate = max(1, min(28, int(row.get(f"billing_date_{i}", 1) or 1)))
                except (ValueError, TypeError):
                    bdate = 1
                try:
                    disc = float(row.get(f"discount_{i}", 0) or 0)
                except (ValueError, TypeError):
                    disc = 0.0
                plans.append({
                    "plan_id": plan["id"],
                    "plan_name": plan["name"],
                    "billing_date": bdate,
                    "discount": disc,
                    "status": "active",
                })

            if plan_errors:
                errors.append({"row": idx, "name": name, "reason": "; ".join(plan_errors)})
                continue

            if not plans:
                errors.append({"row": idx, "name": name, "reason": "At least one plan (plan_name_1) is required"})
                continue

            # Check duplicate WhatsApp
            dup = await db.subscribers.find_one(
                {"whatsapp_number": whatsapp, "operator_id": operator_id, "deleted_at": None}
            )
            if dup:
                skipped.append({"row": idx, "name": name, "reason": f"WhatsApp {whatsapp} already exists"})
                continue

            subscriber = {
                "id": generate_id(),
                "name": name,
                "whatsapp_number": whatsapp,
                "email": row.get("email", "") or None,
                "address": row.get("address", "") or None,
                "plans": plans,
                "status": "active",
                "operator_id": operator_id,
                "created_at": now.isoformat(),
                "updated_at": now.isoformat(),
                "deleted_at": None,
            }
            await db.subscribers.insert_one(subscriber)
            created.append(name)

        return {
            "message": f"Bulk upload complete: {len(created)} created, {len(skipped)} skipped, {len(errors)} errors",
            "created": len(created),
            "skipped": len(skipped),
            "errors": errors[:20],
        }

    @staticmethod
    async def _run_bulk_invoices(job: dict) -> dict:
        """Handler for bulk_upload_invoices."""
        from routers.wallet import deduct_wallet_for_invoice

        payload = job["payload"]
        operator_id = job["operator_id"]
        file_content = payload.get("file_content")
        filename = payload.get("filename", "invoices")

        if isinstance(file_content, str):
            file_content = file_content.encode("utf-8")

        rows = []

        # Parse file
        if filename.lower().endswith((".xlsx", ".xls")):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_content))
            ws = wb.active
            headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append({headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
        else:
            reader = csv.DictReader(io.StringIO(file_content.decode("utf-8-sig")))
            for row in reader:
                rows.append({(k or "").strip().lower(): (v or "").strip() for k, v in row.items()})

        # Pre-load subscribers and plans
        subscribers = await db.subscribers.find(
            {"operator_id": operator_id, "deleted_at": None},
            {"_id": 0, "id": 1, "name": 1, "whatsapp_number": 1},
        ).to_list(5000)
        subscriber_map = {
            (s.get("whatsapp_number") or "").strip(): s
            for s in subscribers
            if (s.get("whatsapp_number") or "").strip()
        }
        plans = await db.operator_plans.find(
            {"operator_id": operator_id, "deleted_at": None}, {"_id": 0}
        ).to_list(1000)
        plan_map = {p["name"].strip().lower(): p for p in plans}

        now = datetime.now(timezone.utc)
        created, skipped, errors = [], [], []

        for idx, row in enumerate(rows, start=2):
            try:
                whatsapp_number = (row.get("subscriber_whatsapp_number") or row.get("whatsapp_number") or "").strip()
                plan_name = (row.get("plan_name") or "").strip().lower()

                if not whatsapp_number or not plan_name:
                    raise ValueError("subscriber_whatsapp_number and plan_name are required")

                subscriber = subscriber_map.get(whatsapp_number)
                if not subscriber:
                    raise ValueError(f"Subscriber with WhatsApp {whatsapp_number} not found")

                plan = plan_map.get(plan_name)
                if not plan:
                    raise ValueError(f"Plan '{row.get('plan_name', '')}' not found")

                from routers.operator import _parse_bulk_invoice_date, _build_invoice_payload
                from models import InvoiceCreate

                service_start_date = _parse_bulk_invoice_date(row.get("service_start_date", ""), "service_start_date")
                service_end_date = _parse_bulk_invoice_date(row.get("service_end_date", ""), "service_end_date")
                due_date = _parse_bulk_invoice_date(row.get("due_date", ""), "due_date")

                if service_end_date <= service_start_date:
                    raise ValueError("service_end_date must be after service_start_date")

                base_amount_raw = (row.get("base_amount") or "").strip()
                discount_raw = (row.get("discount") or "").strip()
                base_amount = float(base_amount_raw) if base_amount_raw else float(plan.get("price", 0))
                discount = float(discount_raw) if discount_raw else 0.0

                if base_amount <= 0:
                    raise ValueError("base_amount must be greater than 0")
                if discount < 0 or discount > base_amount:
                    raise ValueError("discount must be between 0 and base_amount")

                payload_data = await _build_invoice_payload(
                    operator_id,
                    InvoiceCreate(
                        subscriber_id=subscriber["id"],
                        due_date=due_date,
                        line_items=[
                            {
                                "plan_id": plan["id"],
                                "base_amount": base_amount,
                                "discount": discount,
                                "service_start_date": service_start_date,
                                "service_end_date": service_end_date,
                            }
                        ],
                    ),
                )

                invoice = {
                    "id": generate_id(),
                    "invoice_number": await generate_invoice_number_atomic(db),
                    "subscriber_id": subscriber["id"],
                    "subscriber_name": subscriber["name"],
                    "line_items": payload_data["line_items"],
                    "base_amount": payload_data["base_amount"],
                    "discount": payload_data["discount"],
                    "tax_amount": payload_data["tax_amount"],
                    "final_amount": payload_data["final_amount"],
                    "due_date": payload_data["due_date"],
                    "status": "pending",
                    "payment_id": None,
                    "operator_id": operator_id,
                    "created_at": now.isoformat(),
                    "updated_at": now.isoformat(),
                    "deleted_at": None,
                }
                await db.invoices.insert_one(invoice)

                try:
                    await deduct_wallet_for_invoice(operator_id, invoice["id"])
                except Exception as wallet_error:
                    logger.warning(f"Wallet deduction failed for invoice {invoice['id']}: {wallet_error}")

                created.append(invoice["invoice_number"])

            except ValueError as exc:
                errors.append({"row": idx, "reason": str(exc)})
            except Exception as exc:
                errors.append({"row": idx, "reason": f"Unexpected error: {exc}"})

        return {
            "message": f"Bulk upload complete: {len(created)} created, {len(skipped)} skipped, {len(errors)} errors",
            "created": len(created),
            "skipped": len(skipped),
            "errors": errors[:20],
        }

    @staticmethod
    async def _run_bulk_notification(job: dict) -> dict:
        """Handler for bulk_notification."""
        from services.whatsapp_service import WhatsAppService
        from routers.operator import _get_platform_whatsapp_config, _get_whatsapp_template_settings

        payload = job["payload"]
        operator_id = job["operator_id"]
        subscriber_ids = payload.get("subscriber_ids", [])

        wa_config = await _get_platform_whatsapp_config()
        if not wa_config:
            raise ValueError("WhatsApp not configured. Please contact admin.")

        template_settings = await _get_whatsapp_template_settings()
        invoice_template = template_settings.get("invoice_template") or "invoice_notification"

        results = {"sent": 0, "failed": 0, "errors": []}
        wa_service = WhatsAppService(wa_config["phone_number_id"], wa_config["access_token"])

        for subscriber_id in subscriber_ids:
            try:
                subscriber = await db.subscribers.find_one(
                    {"id": subscriber_id, "operator_id": operator_id, "deleted_at": None},
                    {"_id": 0},
                )
                if not subscriber:
                    continue

                invoice = await db.invoices.find_one(
                    {"subscriber_id": subscriber_id, "status": {"$in": ["pending", "overdue"]}, "deleted_at": None},
                    {"_id": 0},
                )
                if invoice:
                    await wa_service.send_invoice_notification(
                        recipient_phone=subscriber["whatsapp_number"],
                        customer_name=subscriber["name"],
                        invoice_number=invoice["invoice_number"],
                        amount=f"₹{invoice['final_amount']:,.2f}",
                        due_date=datetime.fromisoformat(
                            invoice["due_date"].replace("Z", "+00:00")
                        ).strftime("%d %b %Y"),
                        payment_link=invoice.get("payment_link"),
                        template_name_override=invoice_template,
                    )
                    results["sent"] += 1
            except Exception as e:
                results["failed"] += 1
                results["errors"].append({"subscriber_id": subscriber_id, "error": str(e)})

        return results
